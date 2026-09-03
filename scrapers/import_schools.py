"""Import an official CEHRD/IEMIS school CSV or XLSX export into Supabase.

Dry-run is the default. Nothing is written unless --commit is supplied.

Examples:
  python import_schools.py ~/Downloads/SchoolDetails.xlsx --source-url https://cehrd.gov.np/... 
  python import_schools.py schools.csv --source-url https://cehrd.gov.np/... --commit
"""

from __future__ import annotations

import argparse
import csv
import os
import re
import sys
import unicodedata
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from dotenv import load_dotenv
from supabase import create_client

load_dotenv()

ALIASES = {
    "iemis_code": ["iemiscode", "iemisid", "schoolid", "schoolcode", "emiscode"],
    "name": ["schoolname", "nameofschool", "institutionname", "name"],
    "province": ["province", "provincename"],
    "district": ["district", "districtname"],
    "local_level": ["locallevel", "municipality", "localgovernment", "palika"],
    "ward_number": ["ward", "wardno", "wardnumber"],
    "location": ["location", "tole", "address"],
    "ownership_type": ["schooltype", "ownership", "managementtype", "type"],
    "phone": ["phone", "phonenumber", "contact", "mobilenumber"],
    "email": ["email", "emailaddress"],
    "student_count": ["totalstudents", "studenttotal", "enrollment", "totalenrollment"],
    "teacher_count": ["totalteachers", "teachertotal", "teachers"],
    "grades_to": ["highestgrade", "classto", "gradeto", "level"],
}

PROVINCES = {
    "province1": "Koshi", "province01": "Koshi", "koshi": "Koshi",
    "province2": "Madhesh", "madhesh": "Madhesh",
    "province3": "Bagmati", "bagmati": "Bagmati",
    "province4": "Gandaki", "gandaki": "Gandaki",
    "province5": "Lumbini", "lumbini": "Lumbini",
    "province6": "Karnali", "karnali": "Karnali",
    "province7": "Sudurpashchim", "sudurpashchim": "Sudurpashchim", "farwestern": "Sudurpashchim",
}


def normalized(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").strip().lower())


def slugify(value: str) -> str:
    value = unicodedata.normalize("NFKD", value).encode("ascii", "ignore").decode("ascii")
    value = re.sub(r"[^a-zA-Z0-9]+", "-", value.lower()).strip("-")
    return value[:170] or "school"


def integer(value: Any) -> int | None:
    try:
        return int(float(str(value).replace(",", "").strip()))
    except (TypeError, ValueError):
        return None


def rows_from_file(path: Path) -> Iterable[dict[str, Any]]:
    if path.suffix.lower() == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            yield from csv.DictReader(handle)
        return
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Input must be .csv, .xlsx, or .xlsm")
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values = sheet.iter_rows(values_only=True)
    headers = [str(v or "").strip() for v in next(values)]
    for row in values:
        yield dict(zip(headers, row))


def field_map(headers: Iterable[str]) -> dict[str, str]:
    available = {normalized(header): header for header in headers}
    result = {}
    for field, aliases in ALIASES.items():
        for alias in aliases:
            if alias in available:
                result[field] = available[alias]
                break
    return result


def ownership(raw: Any) -> str:
    value = normalized(raw)
    if any(word in value for word in ("community", "samudayik", "government")):
        return "community"
    if any(word in value for word in ("institutional", "private", "boarding")):
        return "institutional"
    if "relig" in value:
        return "religious"
    return "other"


def school_level(highest_grade: int | None) -> str:
    if highest_grade is None:
        return "multiple"
    if highest_grade <= 5:
        return "basic"
    if highest_grade <= 8:
        return "basic"
    if highest_grade <= 10:
        return "secondary"
    return "higher_secondary"


def transform(row: dict[str, Any], mapping: dict[str, str], source_url: str, source_name: str) -> dict[str, Any] | None:
    def get(field: str) -> Any:
        return row.get(mapping.get(field, ""))

    name = str(get("name") or "").strip()
    province_raw = str(get("province") or "").strip()
    district = str(get("district") or "").strip()
    code = str(get("iemis_code") or "").strip().removesuffix(".0")
    if not name or not province_raw or not district:
        return None
    grade_to = integer(get("grades_to"))
    province = PROVINCES.get(normalized(province_raw), province_raw.title())
    suffix = code or normalized(district)[:20]
    return {
        "iemis_code": code or None,
        "name": name,
        "slug": f"{slugify(name)}-{slugify(suffix)}",
        "ownership_type": ownership(get("ownership_type")),
        "school_level": school_level(grade_to),
        "grades_from": 0,
        "grades_to": grade_to,
        "province": province,
        "district": district.title(),
        "local_level": str(get("local_level") or "").strip() or None,
        "ward_number": integer(get("ward_number")),
        "location": str(get("location") or "").strip() or None,
        "phone": str(get("phone") or "").strip() or None,
        "email": str(get("email") or "").strip().lower() or None,
        "student_count": integer(get("student_count")),
        "teacher_count": integer(get("teacher_count")),
        "status": "active",
        "verification_status": "source_verified",
        "source_name": source_name,
        "source_url": source_url,
        "last_verified_at": datetime.now(timezone.utc).isoformat(),
    }


def chunks(items: list[dict[str, Any]], size: int = 250):
    for index in range(0, len(items), size):
        yield items[index:index + size]


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("file", type=Path)
    parser.add_argument("--source-url", required=True, help="Official URL where the dataset was published")
    parser.add_argument("--source-name", default="CEHRD / IEMIS")
    parser.add_argument("--commit", action="store_true", help="Write validated rows to Supabase")
    args = parser.parse_args()

    raw_rows = list(rows_from_file(args.file))
    if not raw_rows:
        print("No rows found", file=sys.stderr)
        return 1
    mapping = field_map(raw_rows[0].keys())
    missing = {"name", "province", "district"} - mapping.keys()
    if missing:
        print(f"Missing required columns: {', '.join(sorted(missing))}", file=sys.stderr)
        print(f"Detected columns: {', '.join(raw_rows[0].keys())}", file=sys.stderr)
        return 1

    records = [record for row in raw_rows if (record := transform(row, mapping, args.source_url, args.source_name))]
    print(f"Validated {len(records)} of {len(raw_rows)} rows")
    print(f"Column mapping: {mapping}")
    if not args.commit:
        print("Dry run complete. Add --commit to upsert these rows.")
        return 0

    url = os.getenv("SUPABASE_URL")
    key = os.getenv("SUPABASE_SERVICE_ROLE_KEY")
    if not url or not key:
        print("SUPABASE_URL and SUPABASE_SERVICE_ROLE_KEY are required for --commit", file=sys.stderr)
        return 1
    client = create_client(url, key)
    conflict_key = "iemis_code" if all(record["iemis_code"] for record in records) else "slug"
    written = 0
    for batch in chunks(records):
        client.table("schools").upsert(batch, on_conflict=conflict_key).execute()
        written += len(batch)
        print(f"Upserted {written}/{len(records)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

